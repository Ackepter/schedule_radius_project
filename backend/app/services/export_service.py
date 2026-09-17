"""Экспорт недельного расписания в PNG/PDF/Excel."""

import io
from datetime import time
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from app.core.config import get_settings
from app.models.entities import LessonTypeEnum, ScheduledLesson, Student
from app.services.pricing_service import calculate_lesson_price

plt.rcParams["font.family"] = "DejaVu Sans"

DAYS_RU = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]
HOUR_START = 9
HOUR_END = 20


def _load_lessons(db: Session, schedule_id: int) -> list[ScheduledLesson]:
    from sqlalchemy import select

    return list(
        db.execute(
            select(ScheduledLesson)
            .options(
                selectinload(ScheduledLesson.lesson_request),
                selectinload(ScheduledLesson.student),
                selectinload(ScheduledLesson.participants),
                selectinload(ScheduledLesson.teacher),
                selectinload(ScheduledLesson.room),
            )
            .where(ScheduledLesson.schedule_id == schedule_id)
        ).scalars()
    )


def _lesson_title(sl: ScheduledLesson, db: Session) -> str:
    if sl.lesson_request and sl.lesson_request.subject:
        return sl.lesson_request.subject.name
    return "Занятие"


def _lesson_type_ru(sl: ScheduledLesson) -> str:
    from app.models.entities import LessonTypeEnum

    return "Групповое" if sl.lesson_type == LessonTypeEnum.group else "Индивидуальное"


def _cost_str(sl: ScheduledLesson, db: Session) -> str:
    if sl.lesson_type == LessonTypeEnum.group:
        subject = sl.lesson_request.subject if sl.lesson_request else None
        if subject:
            n = max(len(sl.participants), 1)
            price = calculate_lesson_price(db, subject.id, "group", n)
            return f"{price} ₽/чел · {int(price * n)} ₽"
    if sl.lesson_request and sl.lesson_request.subject:
        price = calculate_lesson_price(
            db, sl.lesson_request.subject_id, "individual", 1
        )
        return f"{price} ₽"
    return ""


def _fmt(t: time) -> str:
    return f"{t.hour:02d}:{t.minute:02d}"


def render_week_table(
    sl_list: list[ScheduledLesson], db: Session, title: str = "Расписание занятий"
):
    """Возвращает сетку занятий по дням для отрисовки."""
    grid = {d: [] for d in range(7)}
    for sl in sl_list:
        grid[sl.day_of_week].append(sl)
    # раскраска
    colors = [
        "#4c8bf5", "#e8873e", "#3fae69", "#c05bd3", "#d3b13f",
        "#4fc3c3", "#e06a6a", "#7f8fa6", "#9b59b6", "#1abc9c",
    ]
    color_map = {}

    def color_for(sl: ScheduledLesson) -> str:
        key = sl.id
        if key not in color_map:
            color_map[key] = colors[len(color_map) % len(colors)]
        return color_map[key]

    minute_scale = 1.2  # px per minute
    day_width = 148
    rows = []

    for d in range(7):
        cell_lessons = sorted(grid[d], key=lambda x: (x.start_time.hour, x.start_time.minute))
        rows.append(
            {
                "day": DAYS_RU[d],
                "lessons": [
                    {
                        "top": (sl.start_time.hour - HOUR_START) * 60 * minute_scale
                        + sl.start_time.minute * minute_scale,
                        "height": max(
                            22, (sl.end_time.hour * 60 + sl.end_time.minute
                                 - sl.start_time.hour * 60 - sl.start_time.minute)
                            * minute_scale
                        ),
                        "start": _fmt(sl.start_time),
                        "end": _fmt(sl.end_time),
                        "title": _lesson_title(sl, db),
                        "type": _lesson_type_ru(sl),
                        "teacher": f"{sl.teacher.last_name} {sl.teacher.first_name}".strip()
                        if sl.teacher else "—",
                        "room": sl.room.name if sl.room else "—",
                        "color": color_for(sl),
                    }
                    for sl in cell_lessons
                ],
            }
        )
    return rows, minute_scale, day_width


def export_png(db: Session, schedule_id: int) -> bytes:
    sl_list = _load_lessons(db, schedule_id)
    rows, minute_scale, day_width = render_week_table(sl_list, db)

    n_days = 7
    hours = HOUR_END - HOUR_START
    height_per_hour = 60 * minute_scale
    header_h = 46
    total_h = header_h + height_per_hour * hours + 30
    total_w = day_width * n_days + 80

    fig, ax = plt.subplots(figsize=(total_w / 96, total_h / 96), dpi=96)
    ax.set_xlim(0, total_w)
    ax.set_ylim(0, total_h)
    ax.set_aspect("auto")
    ax.axis("off")

    # сетка часов
    for h in range(hours + 1):
        y = header_h + h * height_per_hour
        ax.plot([0, total_w], [y, y], color="#e0e0e0", lw=0.8, zorder=1)
        ax.text(6, y + 4, f"{HOUR_START + h:02d}:00", fontsize=8, ha="left", va="top", color="#555")
        for d in range(n_days):
            x = 80 + d * day_width
            ax.plot([x, x], [header_h, total_h], color="#d0d0d0", lw=1, zorder=1)

    ax.text(total_w / 2, header_h - 10, "НЕДЕЛЬНОЕ РАСПИСАНИЕ", fontsize=14, fontweight="bold",
            ha="center", va="center", color="#222")

    for ri, row in enumerate(rows):
        x0 = 80 + ri * day_width
        ax.text(x0 + day_width / 2, header_h - 28, row["day"], fontsize=11, fontweight="bold",
                ha="center", va="center", color="#333")
        for l in row["lessons"]:
            top = header_h + l["top"]
            ax.add_patch(Rectangle((x0 + 3, top), day_width - 6, l["height"], facecolor=l["color"],
                                   edgecolor="black", lw=0.5, alpha=0.88, zorder=2))
            tx = x0 + 8
            ty = top + 6
            ax.text(tx, ty, f"{l['start']}–{l['end']}", fontsize=7.5, ha="left", va="top",
                    color="white", fontweight="bold", zorder=3)
            ax.text(tx, ty + 12, l["title"], fontsize=8, ha="left", va="top", color="white", zorder=3)
            ax.text(tx, ty + 23, l["type"], fontsize=7, ha="left", va="top", color="#f8f8f8", zorder=3)
            ax.text(tx, ty + 34, l["teacher"], fontsize=7, ha="left", va="top", color="#f8f8f8", zorder=3)
            ax.text(tx, ty + 45, l["room"], fontsize=7, ha="left", va="top", color="#f8f8f8", zorder=3)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=96, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def export_pdf(db: Session, schedule_id: int) -> bytes:
    png = export_png(db, schedule_id)
    from matplotlib.backends.backend_pdf import PdfPages
    from PIL import Image

    buf = io.BytesIO()
    with PdfPages(buf) as pdf:
        img = Image.open(io.BytesIO(png))
        pdf.savefig(__tmp_plot_from_png(img), bbox_inches="tight")
    buf.seek(0)
    return buf.getvalue()


def __tmp_plot_from_png(img):
    import numpy as np

    fig, ax = plt.subplots(figsize=(img.width / 96, img.height / 96))
    ax.imshow(np.asarray(img), aspect="auto")
    ax.axis("off")
    return fig


def export_xlsx(db: Session, schedule_id: int) -> bytes:
    sl_list = _load_lessons(db, schedule_id)
    grid = {d: [] for d in range(7)}
    for sl in sl_list:
        grid[sl.day_of_week].append(sl)

    wb = Workbook()
    ws = wb.active
    ws.title = "Расписание"

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    c = ws.cell(row=1, column=1, value="НЕДЕЛЬНОЕ РАСПИСАНИЕ")
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal="center")

    for idx, day in enumerate(DAYS_RU):
        col = idx + 1
        cell = ws.cell(row=2, column=col, value=day)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    day_cols = [1, 2, 3, 4, 5, 6, 7]
    r = row = 3
    for d in range(7):
        lessons = sorted(grid.get(d, []), key=lambda x: (x.start_time.hour, x.start_time.minute))
        col = day_cols[d]
        for sl in lessons:
            cell = ws.cell(row=row, column=col)
            lines = [
                f"{_fmt(sl.start_time)}-{_fmt(sl.end_time)}",
                _lesson_title(sl, db),
                _lesson_type_ru(sl),
                f"{sl.teacher.last_name} {sl.teacher.first_name}".strip() if sl.teacher else "—",
                sl.room.name if sl.room else "—",
            ]
            if sl.lesson_type == LessonTypeEnum.group:
                n = max(len(sl.participants), 1)
                lines.append(f"{n} чел.")
            else:
                lines.append(sl.student.full_name if sl.student else "")
            cell.value = "\n".join(lines)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            row += 1
        if lessons:
            row += 1
        else:
            ws.cell(row=row, column=col, value="—").border = border

    width = 24
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_export(db: Session, schedule_id: int, format, file_path: Path) -> None:
    from app.schemas.schemas import ExportFormat

    if format == ExportFormat.png:
        data = export_png(db, schedule_id)
    elif format == ExportFormat.pdf:
        data = export_pdf(db, schedule_id)
    else:
        data = export_xlsx(db, schedule_id)
    file_path.write_bytes(data)